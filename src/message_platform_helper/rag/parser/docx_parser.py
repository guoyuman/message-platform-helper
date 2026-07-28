"""Word document parser."""

from __future__ import annotations

import re
from pathlib import Path

from ..models import Document, DocumentSection, SectionKind


HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
TABLE_RE = re.compile(r"^\s*\|.*\|\s*$")


class DocxParser:
    def parse(self, document: Document) -> list[DocumentSection]:
        source = Path(document.metadata.source)
        if source.exists() and source.suffix.lower() == ".docx":
            parsed = _parse_docx_file(source, document)
            if parsed:
                return parsed
        return _parse_text(document)


def _parse_docx_file(path: Path, document: Document) -> list[DocumentSection]:
    try:
        from docx import Document as DocxDocument  # type: ignore[import-untyped]
    except ImportError:
        return []

    docx = DocxDocument(path)
    sections: list[DocumentSection] = []
    current_title = document.title
    current_level = 0
    buffer: list[str] = []

    def flush(kind: SectionKind = "paragraph") -> None:
        nonlocal buffer
        content = "\n\n".join(buffer).strip()
        if content:
            sections.append(DocumentSection(title=current_title, level=current_level, content=content, kind=kind))
        buffer = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.lower().startswith("heading"):
            flush()
            current_title = text
            current_level = _heading_level(style_name)
            continue
        buffer.append(text)
    flush()

    for table in docx.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        content = "\n".join(row for row in rows if row.strip())
        if content.strip():
            sections.append(DocumentSection(title=current_title, level=current_level, content=content, kind="table"))
    return sections


def _parse_text(document: Document) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    current_title = document.title
    current_level = 0
    buffer: list[str] = []

    def flush(kind: SectionKind = "paragraph") -> None:
        nonlocal buffer
        content = "\n".join(buffer).strip()
        if content:
            sections.append(DocumentSection(title=current_title, level=current_level, content=content, kind=kind))
        buffer = []

    for block in _blocks(document.content):
        heading = HEADING_RE.match(block)
        if heading:
            flush()
            current_level = len(heading.group(1))
            current_title = heading.group(2).strip()
            continue
        kind: SectionKind = "table" if all(TABLE_RE.match(line) for line in block.splitlines() if line.strip()) else "paragraph"
        if buffer and kind != _text_block_kind("\n".join(buffer)):
            flush(kind=_text_block_kind("\n".join(buffer)))
        buffer.append(block)
    flush(kind=_text_block_kind("\n".join(buffer)))
    if not sections and document.content.strip():
        sections.append(DocumentSection(title=document.title, level=0, content=document.content.strip(), kind="text"))
    return sections


def _blocks(text: str) -> list[str]:
    return [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]


def _text_block_kind(block: str) -> SectionKind:
    return "table" if all(TABLE_RE.match(line) for line in block.splitlines() if line.strip()) else "paragraph"


def _heading_level(style_name: str) -> int:
    match = re.search(r"(\d+)", style_name)
    if not match:
        return 1
    return max(1, int(match.group(1)))
