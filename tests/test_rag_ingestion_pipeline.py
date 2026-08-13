from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from message_platform_helper.rag import (
    Document,
    DocumentMetadata,
    DocumentSection,
    MarkdownChunkStrategy,
    ParentChildChunkStrategy,
    RecursiveChunkStrategy,
)
from message_platform_helper.rag import MarkdownLoader, MarkdownParser
from message_platform_helper.rag.loader.docx_loader import read_docx_text
from message_platform_helper.rag.parser.docx_parser import DocxParser
from message_platform_helper.rag.parser.pdf_parser import PdfParser
from message_platform_helper.rag import XlsxLoader

from tests.fakes import InMemoryKnowledgeBase


class RagIngestionPipelineTests(unittest.TestCase):
    def test_markdown_loader_returns_document_with_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "email_failure.md"
            path.write_text(
                "\n".join(
                    [
                        "---",
                        "title: Email failure",
                        "tags: mail, ops",
                        "---",
                        "# Email failure",
                        "",
                        "Check delivery logs.",
                    ]
                ),
                encoding="utf-8",
            )

            document = MarkdownLoader().load(path)

        self.assertEqual(document.title, "Email failure")
        self.assertEqual(document.metadata.format, "markdown")
        self.assertIn("mail", document.metadata.tags)
        self.assertIn("Check delivery logs.", document.content)

    def test_markdown_parser_preserves_heading_context(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "guide.md"
            path.write_text("# Mail\n\nIntro.\n\n## Troubleshooting\n\nKeep paragraphs whole.", encoding="utf-8")
            parsed = MarkdownParser().parse(MarkdownLoader().load(path))

        self.assertEqual([section.title for section in parsed], ["Mail", "Troubleshooting"])
        self.assertEqual(parsed[1].level, 2)
        self.assertEqual(parsed[1].content, "Keep paragraphs whole.")

    def test_knowledge_base_ingests_chunks_with_parent_child_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "guide.md"
            path.write_text("# Mail\n\nIntro.\n\n## Troubleshooting\n\nCheck logs.", encoding="utf-8")
            kb = InMemoryKnowledgeBase()

            chunks = kb.ingest_file(path, tags=["mail"], strategy=MarkdownChunkStrategy(embedding_fn=_topic_embedding))

        self.assertEqual(len(chunks), 2)
        self.assertEqual(kb.count(), 2)
        self.assertTrue(any(chunk.title == "Troubleshooting" for chunk in chunks))
        self.assertTrue(all(chunk.source.endswith("guide.md") for chunk in chunks))
        self.assertTrue(all(chunk.tags == ["mail"] for chunk in chunks))

    def test_markdown_chunking_uses_paragraph_boundaries_before_size_limit(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "guide.md"
            path.write_text("# Mail\n\nIntro one.\n\nIntro two.\n\n## Troubleshooting\n\nCheck logs.", encoding="utf-8")
            kb = InMemoryKnowledgeBase()

            chunks = kb.ingest_file(path, tags=["mail"], strategy=MarkdownChunkStrategy())

        self.assertEqual([chunk.title for chunk in chunks], ["Mail", "Mail", "Troubleshooting"])
        self.assertEqual([chunk.content for chunk in chunks], ["Intro one.", "Intro two.", "Check logs."])

    def test_recursive_chunker_keeps_plain_paragraph_under_limit(self) -> None:
        chunks = _chunk_text("普通段落内容。", chunk_size=20, max_chunk_size=20)

        self.assertEqual([chunk.content for chunk in chunks], ["普通段落内容。"])

    def test_recursive_chunker_semantically_merges_related_candidate_paragraphs(self) -> None:
        chunks = _chunk_text(
            "MySQL supports replication.\n\nMySQL replicas handle reads.",
            chunk_size=80,
            max_chunk_size=80,
            embedding_fn=_topic_embedding,
        )

        self.assertEqual(len(chunks), 1)
        self.assertIn("replication.\n\nMySQL replicas", chunks[0].content)

    def test_recursive_chunker_keeps_unrelated_candidate_paragraphs_separate(self) -> None:
        chunks = _chunk_text(
            "MySQL supports replication.\n\nInvoice templates render email content.",
            chunk_size=80,
            max_chunk_size=80,
            embedding_fn=_topic_embedding,
        )

        self.assertEqual([chunk.content for chunk in chunks], ["MySQL supports replication.", "Invoice templates render email content."])

    def test_recursive_chunker_splits_long_sentence_without_period(self) -> None:
        chunks = _chunk_text("超长句子" * 30, chunk_size=25, max_chunk_size=25)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.content) <= 25 for chunk in chunks))

    def test_recursive_chunker_prefers_semicolon_before_lower_priority_splits(self) -> None:
        chunks = _chunk_text("alpha beta gamma;delta epsilon zeta;theta iota kappa", chunk_size=22, max_chunk_size=22)

        self.assertEqual([chunk.content for chunk in chunks], ["alpha beta gamma;", "delta epsilon zeta;", "theta iota kappa"])

    def test_recursive_chunker_falls_back_to_comma(self) -> None:
        chunks = _chunk_text("alpha beta gamma,delta epsilon zeta,theta iota kappa", chunk_size=22, max_chunk_size=22)

        self.assertEqual([chunk.content for chunk in chunks], ["alpha beta gamma,", "delta epsilon zeta,", "theta iota kappa"])

    def test_recursive_chunker_hard_splits_text_without_separators(self) -> None:
        chunks = _chunk_text("x" * 53, chunk_size=20, max_chunk_size=20)

        self.assertEqual([len(chunk.content) for chunk in chunks], [20, 20, 13])

    def test_recursive_chunker_does_not_split_images(self) -> None:
        document = Document(id="doc-image", title="Image", content="", metadata=DocumentMetadata(format="pdf"))
        section = DocumentSection(title="Image", level=1, kind="image", content="[Image 1.1: page=1, size=20x30]")

        chunks = RecursiveChunkStrategy(chunk_size=10, max_chunk_size=10).chunk(document, [section])

        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].content, section.content)

    def test_recursive_chunker_ignores_accidental_blank_line_when_semantically_related(self) -> None:
        chunks = _chunk_text(
            "MySQL 支持主从复制。\n\n主库负责写入，从库负责读取。",
            chunk_size=60,
            max_chunk_size=60,
            embedding_fn=_topic_embedding,
        )

        self.assertEqual(len(chunks), 1)
        self.assertIn("\n\n", chunks[0].content)

    def test_recursive_chunker_does_not_merge_across_headings(self) -> None:
        document = Document(
            id="doc-heading",
            title="Guide",
            content="",
            metadata=DocumentMetadata(format="markdown"),
        )
        sections = [
            DocumentSection(title="数据库配置", level=2, kind="paragraph", content="MySQL supports replication."),
            DocumentSection(title="模板配置", level=2, kind="paragraph", content="MySQL replicas handle reads."),
        ]

        chunks = RecursiveChunkStrategy(chunk_size=80, max_chunk_size=80, embedding_fn=_topic_embedding).chunk(document, sections)

        self.assertEqual([chunk.title for chunk in chunks], ["数据库配置", "模板配置"])

    def test_recursive_chunker_respects_max_size_even_when_similarity_is_high(self) -> None:
        chunks = _chunk_text(
            "MySQL " + ("replication " * 6) + "\n\nMySQL " + ("replica " * 6),
            chunk_size=50,
            max_chunk_size=50,
            embedding_fn=_topic_embedding,
        )

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.content) <= 50 for chunk in chunks))

    def test_recursive_chunker_adds_bounded_overlap(self) -> None:
        chunks = _chunk_text("First sentence. Second sentence. Third sentence.", chunk_size=20, max_chunk_size=36, chunk_overlap=16)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(chunks[1].content.startswith("Second sentence."))
        self.assertTrue(all(len(chunk.content) <= 36 for chunk in chunks))

    def test_xlsx_loader_extracts_sheet_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "roles.xlsx"
            with ZipFile(path, "w") as workbook:
                workbook.writestr(
                    "xl/workbook.xml",
                    """<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                    <sheets><sheet name="岗位表" sheetId="1" r:id="rId1"/></sheets></workbook>""",
                )
                workbook.writestr(
                    "xl/_rels/workbook.xml.rels",
                    """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                    <Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>""",
                )
                workbook.writestr(
                    "xl/worksheets/sheet1.xml",
                    """<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
                    <sheetData><row><c t="inlineStr"><is><t>单位</t></is></c>
                    <c t="inlineStr"><is><t>岗位</t></is></c></row></sheetData></worksheet>""",
                )

            document = XlsxLoader().load(path)

        self.assertEqual(document.metadata.format, "xlsx")
        self.assertIn("# 岗位表", document.content)
        self.assertIn("单位 | 岗位", document.content)

    def test_parent_child_chunker_keeps_semantic_table_rows_together(self) -> None:
        document = Document(
            id="doc-table",
            title="岗位表",
            content="",
            metadata=DocumentMetadata(format="docx", source="table.docx"),
        )
        section = DocumentSection(
            title="岗位表",
            level=1,
            kind="table",
            content="\n".join(["单位 | 岗位", *[f"单位{i} | 岗位{i}" for i in range(12)]]),
        )

        chunks = ParentChildChunkStrategy(child_chunk_size=70).chunk(document, [section])

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.content.startswith("单位 | 岗位") for chunk in chunks))
        self.assertTrue(all(chunk.metadata["section_kind"] == "table" for chunk in chunks))

    def test_pdf_parser_preserves_page_table_and_image_sections(self) -> None:
        document = Document(
            id="pdf",
            title="Guide",
            content="[Page 2]\n正文\n\n[Table 2.1]\nA | B\n1 | 2\n\n[Image 2.1: page=2, size=20x30]",
            metadata=DocumentMetadata(format="pdf", source="guide.pdf"),
        )

        sections = PdfParser().parse(document)

        self.assertEqual([section.kind for section in sections], ["paragraph", "table", "image"])
        self.assertEqual(sections[1].metadata["page"], 2)
        self.assertEqual(sections[0].title, "Guide")
        self.assertEqual(sections[2].title, "Guide Image 2.1")

    def test_docx_loader_preserves_paragraph_table_and_image_markers_in_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mixed.docx"
            _write_mixed_docx(path)

            content = read_docx_text(path)
            sections = DocxParser().parse(
                Document(id="docx", title="Mixed", content=content, metadata=DocumentMetadata(format="docx", source=str(path)))
            )

        self.assertLess(content.index("Intro"), content.index("A | B"))
        self.assertLess(content.index("A | B"), content.index("[Image 1:"))
        self.assertEqual([section.kind for section in sections], ["paragraph", "table", "image"])


def _write_mixed_docx(path: Path) -> None:
    import base64

    from docx import Document as DocxDocument  # type: ignore[import-untyped]

    image_path = path.with_suffix(".png")
    image_path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )
    )
    document = DocxDocument()
    document.add_paragraph("Intro")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(str(image_path))
    document.save(path)


def _chunk_text(
    text: str,
    *,
    chunk_size: int,
    max_chunk_size: int,
    chunk_overlap: int = 0,
    embedding_fn=None,
):
    document = Document(id="doc-text", title="Text", content=text, metadata=DocumentMetadata(format="text"))
    return RecursiveChunkStrategy(
        chunk_size=chunk_size,
        max_chunk_size=max_chunk_size,
        chunk_overlap=chunk_overlap,
        embedding_fn=embedding_fn,
    ).chunk(document, [])


def _topic_embedding(paragraphs: list[str]) -> list[list[float]]:
    vectors: list[list[float]] = []
    for paragraph in paragraphs:
        if "mysql" in paragraph.lower() or "MySQL" in paragraph or "主" in paragraph or "从" in paragraph:
            vectors.append([1.0, 0.0])
        else:
            vectors.append([0.0, 1.0])
    return vectors


if __name__ == "__main__":
    unittest.main()
